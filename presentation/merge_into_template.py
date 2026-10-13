import copy
import io
import sys

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

TEMPLATE, CONTENT, OUT = sys.argv[1:4]

tpl = Presentation(TEMPLATE)
ids = tpl.slides._sldIdLst
for sld in list(ids):                       # remove the template's own sample slides, keep layouts/master/logo
    tpl.part.drop_rel(sld.rId)
    ids.remove(sld)

layouts = {layout.name: layout for layout in tpl.slide_layouts}
src = Presentation(CONTENT)


def set_text(placeholder, text):
    """Fill a placeholder keeping its template formatting (font, size, colour come from the layout)."""
    placeholder.text_frame.text = text


for number, source in enumerate(src.slides, start=1):
    has_subtitle = any(sh.name == "SLIDE_SUBTITLE" for sh in source.shapes)   # a cover has a subtitle
    layout = layouts["1_Title Slide"] if has_subtitle else layouts["1_Title Only"]
    dest = tpl.slides.add_slide(layout)

    title_text = subtitle_text = None
    for shape in list(source.shapes):
        if shape.name == "SLIDE_TITLE":
            title_text = shape.text_frame.text
        elif shape.name == "SLIDE_SUBTITLE":
            subtitle_text = shape.text_frame.text
    set_text(dest.shapes.title, title_text)
    if subtitle_text is not None:
        set_text(dest.placeholders[1], subtitle_text)

    tree = dest.shapes._spTree
    for shape in source.shapes:
        if shape.name in ("SLIDE_TITLE", "SLIDE_SUBTITLE"):
            continue
        element = copy.deepcopy(shape._element)
        for blip in element.iter(qn("a:blip")):          # pictures: copy the image into the new slide
            old = blip.get(qn("r:embed"))
            _, new = dest.part.get_or_add_image_part(io.BytesIO(source.part.related_part(old).blob))
            blip.set(qn("r:embed"), new)
        tree.append(element)

    for ph in layout.placeholders:                      # slide number: python-pptx does not add it by itself
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            tree.append(copy.deepcopy(ph._element))

    for index, cnv in enumerate(tree.iter(qn("p:cNvPr")), start=2):   # unique shape ids
        cnv.set("id", str(index))

    if source.has_notes_slide:
        dest.notes_slide.notes_text_frame.text = source.notes_slide.notes_text_frame.text

tpl.save(OUT)
print(f"saved {OUT}: {len(tpl.slides)} slides, layouts kept: {len(tpl.slide_layouts)}")
